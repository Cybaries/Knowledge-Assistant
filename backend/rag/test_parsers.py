from pathlib import Path

from django.test import SimpleTestCase
from rag.parsers import (DocumentParsingError, parse_document, parse_docx,
                         parse_txt)


class ParserTests(SimpleTestCase):
    def test_parse_txt(self):
        path = Path("/tmp/test_parser.txt")
        path.write_text("Hello from a text document.", encoding="utf-8")

        try:
            self.assertEqual(
                parse_txt(str(path)),
                "Hello from a text document.",
            )
        finally:
            path.unlink(missing_ok=True)

    def test_parse_docx(self):
        from docx import Document

        path = Path("/tmp/test_parser.docx")
        document = Document()
        document.add_paragraph("First paragraph.")
        document.add_paragraph("Second paragraph.")
        document.save(path)

        try:
            text = parse_docx(str(path))

            self.assertIn("First paragraph.", text)
            self.assertIn("Second paragraph.", text)
        finally:
            path.unlink(missing_ok=True)

    def test_parse_document_dispatches_txt(self):
        path = Path("/tmp/test_dispatch.txt")
        path.write_text("Dispatcher test.", encoding="utf-8")

        try:
            self.assertEqual(
                parse_document(str(path)),
                "Dispatcher test.",
            )
        finally:
            path.unlink(missing_ok=True)

    def test_unsupported_file_type(self):
        with self.assertRaises(DocumentParsingError):
            parse_document("/tmp/test.xyz")

    def test_missing_txt_file(self):
        with self.assertRaises(DocumentParsingError):
            parse_txt("/tmp/does-not-exist.txt")